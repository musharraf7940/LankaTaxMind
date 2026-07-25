from fastapi import APIRouter, HTTPException, Header, Request
from pydantic import BaseModel
from pinecone import Pinecone
from google import genai
from google.genai import types
from firebase_admin import auth as firebase_auth
from slowapi import Limiter
from slowapi.util import get_remote_address
import os
import time
from dotenv import load_dotenv

from fastapi import UploadFile, File
import fitz
from docx import Document as DocxDocument
from docx import Document as DocxWriter
from fastapi.responses import StreamingResponse
import openpyxl
import io


load_dotenv(dotenv_path=os.path.join(os.path.dirname(__file__), '..', '.env'))

router = APIRouter()
limiter = Limiter(key_func=get_remote_address)

PINECONE_API_KEY = os.getenv("PINECONE_API_KEY")
GEMINI_API_KEY = os.getenv("GEMINI_API_KEY")
PINECONE_INDEX = os.getenv("PINECONE_INDEX", "auditmind")

pc = Pinecone(api_key=PINECONE_API_KEY)
index = pc.Index(PINECONE_INDEX)

client = genai.Client(api_key=GEMINI_API_KEY)

from huggingface_hub import InferenceClient

HF_TOKEN = os.getenv("HF_TOKEN")
hf_client = InferenceClient(token=HF_TOKEN)
EMBEDDING_MODEL_NAME = "sentence-transformers/all-MiniLM-L6-v2"

def get_embedding(text: str):
    result = hf_client.feature_extraction(text, model=EMBEDDING_MODEL_NAME)
    return list(result)

PRIMARY_MODEL = "gemini-2.5-flash"
FALLBACK_MODEL = "gemini-2.0-flash"


def generate_with_retry(contents, config=None, max_retries=3):
    models = [PRIMARY_MODEL, FALLBACK_MODEL]
    for model in models:
        for attempt in range(max_retries):
            try:
                if config:
                    response = client.models.generate_content(
                        model=model,
                        contents=contents,
                        config=config
                    )
                else:
                    response = client.models.generate_content(
                        model=model,
                        contents=contents
                    )
                return response, model
            except Exception as e:
                error_str = str(e)
                if "503" in error_str or "UNAVAILABLE" in error_str or "429" in error_str:
                    wait_time = (attempt + 1) * 3
                    print(f"Model {model} unavailable. Retrying in {wait_time}s... (attempt {attempt + 1}/{max_retries})")
                    time.sleep(wait_time)
                    continue
                else:
                    raise e
    raise Exception("All models unavailable after retries.")


def detect_and_translate(question: str):
    detection_prompt = (
        "Detect the language of the following text. It will be English, Sinhala, or Tamil.\n"
        "If it is not English, translate it to English.\n"
        "Respond ONLY in this exact format, nothing else:\n"
        "LANGUAGE: <English/Sinhala/Tamil>\n"
        "TRANSLATED: <the English translation, or the original text if it was already English>\n\n"
        f"TEXT: {question}"
    )

    response, _ = generate_with_retry(contents=detection_prompt)
    result = response.text.strip()

    language = "English"
    translated = question

    for line in result.split("\n"):
        if line.startswith("LANGUAGE:"):
            language = line.replace("LANGUAGE:", "").strip()
        elif line.startswith("TRANSLATED:"):
            translated = line.replace("TRANSLATED:", "").strip()

    return language, translated


def has_draft_signal(text: str) -> bool:
    phrases = ["draft a", "draft the", "draft an", "would you like me to draft", "shall i draft", "i can draft"]
    lowered = text.lower()
    return any(phrase in lowered for phrase in phrases)


class QueryRequest(BaseModel):
    question: str
    history: list = []


class QueryResponse(BaseModel):
    answer: str
    sources: list
    warning: str = ""
    draft_offered: bool = False


@router.post("/ask", response_model=QueryResponse)
@limiter.limit("10/minute")
async def ask_question(request: Request, query: QueryRequest, authorization: str = Header(None)):

    if not authorization or not authorization.startswith("Bearer "):
        raise HTTPException(status_code=401, detail="Missing or invalid authorization token")

    id_token = authorization.split("Bearer ")[1]

    try:
        firebase_auth.verify_id_token(id_token)
    except Exception:
        raise HTTPException(status_code=401, detail="Invalid or expired token")

    if not query.question.strip():
        raise HTTPException(status_code=400, detail="Question cannot be empty")

    casual_greeting_words = {
        "hi", "hello", "hey", "thanks", "thank", "ok", "okay", "good",
        "great", "nice", "perfect", "cool", "awesome", "bye", "goodbye",
        "yes", "no", "sure", "welcome", "appreciate", "helpful", "well done",
        "what's", "whats", "up", "how", "are", "doing", "going"
    }

    tax_keywords = {
        "tax", "vat", "act", "section", "law", "rate", "penalty", "irs", "ird",
        "revenue", "duty", "levy", "exemption", "assessment", "appeal", "case",
        "court", "invoice", "income", "deduction", "compliance", "audit",
        "gazette", "amendment", "provision", "statute", "clause", "surcharge"
    }

    normalized_words = set(query.question.strip().lower().strip(".!?").split())

    is_casual = (
        len(normalized_words) <= 8
        and normalized_words.issubset(casual_greeting_words | {"you", "it", "is", "to", "from", "for", "me", "i", "so", "very", "much", "again", "there"})
        and not (normalized_words & tax_keywords)
    )

    if is_casual:
        history_text = ""
        if query.history:
            recent_history = query.history[-6:]
            for turn in recent_history:
                role = "User" if turn.get("role") == "user" else "LankaTaxMind"
                history_text += f"{role}: {turn.get('content', '')}\n\n"

        casual_response, _ = generate_with_retry(
            contents=(
                "You are LankaTaxMind, a friendly and professional AI assistant for Sri Lankan tax law. "
                "The user sent a short conversational message, not a legal question. "
                "Respond naturally and warmly, staying in character as a tax law assistant, "
                "without adding any disclaimers about web search or knowledge base limitations.\n\n"
                f"RECENT CONVERSATION:\n{history_text}\n"
                f"USER MESSAGE: {query.question}"
            )
        )
        return QueryResponse(
            answer=casual_response.text.strip(),
            sources=[],
            warning=""
        )

    detected_language, english_question = detect_and_translate(query.question)

    query_vector = get_embedding(english_question)
    results = index.query(
        vector=query_vector,
        top_k=5,
        include_metadata=True
    )

    matches = results.get("matches", [])

    has_heading_match = any(
        match.get("metadata", {}).get("type") == "section_heading" and match["score"] >= 0.40
        for match in matches
    )

    if not matches or (matches[0]["score"] < 0.65 and not has_heading_match):
        history_text = ""
        if query.history:
            recent_history = query.history[-6:]
            for turn in recent_history:
                role = "User" if turn.get("role") == "user" else "LankaTaxMind"
                history_text += f"{role}: {turn.get('content', '')}\n\n"

        try:
            web_response, model_used = generate_with_retry(
                contents=(
                    "You are LankaTaxMind, a senior tax partner AI with deep expertise in "
                    "Sri Lankan tax law and legal frameworks.\n\n"
                    "The internal knowledge base does not contain specific references for "
                    "the following query. Search the web for current, reliable legal sources "
                    "and provide a professional answer with citations. "
                    "If the question relates to disputing, objecting to, appealing, or requesting something "
                    "from a tax authority, end your response with a natural one-line offer such as "
                    "'Would you like me to draft a formal [objection letter/appeal/request] based on this?'\n\n"
                    f"RECENT CONVERSATION HISTORY (for context on follow-up questions):\n{history_text}\n"
                    f"QUESTION: {query.question}"
                ),
                config=types.GenerateContentConfig(
                    tools=[types.Tool(google_search=types.GoogleSearch())]
                )
            )

            web_answer = ""
            for part in web_response.candidates[0].content.parts:
                if hasattr(part, "text"):
                    web_answer += part.text

            return QueryResponse(
                answer=web_answer.strip(),
                sources=[],
                warning="Answer sourced from web search — not from uploaded documents. Verify with authoritative sources.",
                draft_offered=has_draft_signal(web_answer)
            )

        except Exception:
            return QueryResponse(
                answer=(
                    "LankaTaxMind is currently experiencing high demand on its AI service. "
                    "Your question has been received but could not be processed at this moment. "
                    "Please try again in 30 seconds. If the issue persists, "
                    "kindly contact your system administrator."
                ),
                sources=[],
                warning="Service temporarily unavailable. Please retry shortly."
            )

    context = ""
    sources = []

    for match in matches:
        metadata = match.get("metadata", {})
        context += f"\n---\nSource: {metadata.get('source', 'Unknown')}\n"
        context += metadata.get("text", "") + "\n"
        sources.append({
            "source": metadata.get("source", "Unknown"),
            "score": round(match["score"], 4),
            "chunk_index": metadata.get("chunk_index", 0),
            "text": metadata.get("text", "")
        })

    history_text = ""
    if query.history:
        recent_history = query.history[-6:]
        for turn in recent_history:
            role = "User" if turn.get("role") == "user" else "LankaTaxMind"
            history_text += f"{role}: {turn.get('content', '')}\n\n"

    prompt = (
        "You are LankaTaxMind, a senior tax partner AI assistant with deep expertise "
        "in statutory interpretation, case law analysis, and regulatory frameworks "
        "for Sri Lankan tax practice.\n\n"
        "STRICT BEHAVIOURAL RULES:\n"
        "1. Before answering, restate the user's question in your own words to demonstrate full understanding.\n"
        "2. Provide a clear, professional answer supported by authoritative citations from the context documents provided below.\n"
        "3. When citing law or cases from the context, include exact references such as 'Section X, page Y' or 'Case Z v. A, [Year] at paragraph/page'.\n"
        "4. If the provided context contains no applicable references, clearly state: 'The uploaded knowledge base does not contain specific references for this query.'\n"
        "5. STRICTLY NO generic responses. Every answer must be grounded within the legal and regulatory framework found in the provided context.\n"
        "6. Structure responses logically with clear headings where appropriate.\n"
        "7. Use professional legal language throughout.\n"
        "8. Explicitly justify every conclusion with citations to law, regulation, case authority or applicable standard found in the context.\n"
        "9. NEVER fabricate citations or generate references not present in the context below.\n"
        "10. After citing the relevant provisions from the context, expand with professional legal reasoning, "
        "practical implications, and related considerations — clearly distinguishing between what the documents "
        "state and your professional analysis.\n"
        "11. When sources conflict, resolve using this hierarchy: legislation over commentary, "
        "consolidated/amended provisions over superseded ones, and later judicial authority over earlier authority. "
        "If the conflict cannot be resolved confidently, present the competing interpretations rather than "
        "asserting one as unquestionably correct.\n"
        f"12. Respond entirely in {detected_language}, regardless of the language of the source documents. "
        "Keep legal citations (Act names, section numbers) in their original English form even when the "
        "surrounding explanation is in Sinhala or Tamil, since these are official statutory references.\n"
        "13. If the user's question relates to disputing, objecting to, appealing, or requesting something "
        "from a tax authority (e.g. an assessment, penalty, or deadline), end your response with a natural "
        "one-line offer such as 'Would you like me to draft a formal [objection letter/appeal/request] based "
        "on this?' Only include this offer when genuinely relevant — not for general informational questions.\n\n"
        f"RECENT CONVERSATION HISTORY (for context on follow-up questions):\n{history_text}\n"
        f"CONTEXT FROM TAX LAW KNOWLEDGE BASE:\n{context}\n\n"
        f"CURRENT USER QUESTION (original): {query.question}\n"
        f"CURRENT USER QUESTION (English, for reference): {english_question}\n\n"
        "Begin your response by restating the question, then provide your structured professional answer with citations and analysis:"
    )

    try:
        response, model_used = generate_with_retry(contents=prompt)
        answer = response.text.strip()

        if "does not contain specific references" in answer.lower() or "knowledge base does not contain" in answer.lower():
            try:
                web_response, _ = generate_with_retry(
                    contents=(
                        "You are LankaTaxMind, a senior tax partner AI with deep expertise in "
                        "Sri Lankan tax law and legal frameworks.\n\n"
                        "The internal knowledge base does not contain specific references for "
                        "the following query. Search the web for current, reliable legal sources "
                        "and provide a professional answer with citations. "
                        "If the question relates to disputing, objecting to, appealing, or requesting something "
                        "from a tax authority, end your response with a natural one-line offer such as "
                        "'Would you like me to draft a formal [objection letter/appeal/request] based on this?'\n\n"
                        f"RECENT CONVERSATION HISTORY (for context on follow-up questions):\n{history_text}\n"
                        f"QUESTION: {query.question}"
                    ),
                    config=types.GenerateContentConfig(
                        tools=[types.Tool(google_search=types.GoogleSearch())]
                    )
                )

                web_answer = ""
                for part in web_response.candidates[0].content.parts:
                    if hasattr(part, "text"):
                        web_answer += part.text

                return QueryResponse(
                    answer=web_answer.strip(),
                    sources=sources,
                    warning="Primary answer supplemented with web search — uploaded documents did not contain a direct reference. Verify with authoritative sources.",
                    draft_offered=has_draft_signal(web_answer)
                )
            except Exception:
                pass

        return QueryResponse(
            answer=answer,
            sources=sources,
            draft_offered=has_draft_signal(answer)
        )

    except Exception:
        return QueryResponse(
            answer=(
                "LankaTaxMind is currently experiencing high demand on its AI service. "
                "Your question has been received but could not be processed at this moment. "
                "Please try again in 30 seconds. If the issue persists, "
                "kindly contact your system administrator."
            ),
            sources=sources,
            warning="Service temporarily unavailable. Please retry shortly."
        )


class ExplainRequest(BaseModel):
    question: str
    chunk_text: str


class ExplainResponse(BaseModel):
    words: list
    scores: list


@router.post("/explain", response_model=ExplainResponse)
@limiter.limit("15/minute")
async def explain_source(request: Request, body: ExplainRequest, authorization: str = Header(None)):

    import numpy as np
    import shap

    if not authorization or not authorization.startswith("Bearer "):
        raise HTTPException(status_code=401, detail="Missing or invalid authorization token")

    id_token = authorization.split("Bearer ")[1]

    try:
        firebase_auth.verify_id_token(id_token)
    except Exception:
        raise HTTPException(status_code=401, detail="Invalid or expired token")

    query_embedding = np.array(get_embedding(body.question))

    def similarity_score(texts):
        embeddings = np.array([get_embedding(t) for t in texts])
        scores = np.dot(embeddings, query_embedding) / (
            np.linalg.norm(embeddings, axis=1) * np.linalg.norm(query_embedding) + 1e-8
        )
        return scores

    masker = shap.maskers.Text(tokenizer=r"\W+")
    explainer = shap.Explainer(similarity_score, masker)

    truncated_text = " ".join(body.chunk_text.split()[:120])

    shap_values = explainer([truncated_text])

    words = shap_values[0].data.tolist()
    scores = shap_values[0].values.tolist()

    stopwords = {"the", "a", "an", "of", "to", "in", "is", "and", "or", "for", "by",
                 "on", "at", "as", "be", "this", "that", "with", "from", "shall"}

    filtered = [
        (w, s) for w, s in zip(words, scores)
        if w.lower().strip() not in stopwords
        and len(w.strip()) > 1
        and any(c.isalnum() for c in w)
    ]

    ranked = sorted(filtered, key=lambda x: abs(x[1]), reverse=True)
    top_words = ranked[:10]

    return ExplainResponse(
        words=[w for w, s in top_words],
        scores=[round(float(s), 4) for w, s in top_words]
    )


def extract_text_from_upload(filename: str, file_bytes: bytes) -> str:
    ext = filename.lower().split(".")[-1]

    if ext == "pdf":
        doc = fitz.open(stream=file_bytes, filetype="pdf")
        text = ""
        for page in doc:
            text += page.get_text()
        doc.close()
        return text.strip()

    elif ext == "docx":
        doc = DocxDocument(io.BytesIO(file_bytes))
        text = ""
        for para in doc.paragraphs:
            if para.text.strip():
                text += para.text.strip() + "\n"
        return text.strip()

    elif ext == "xlsx":
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes))
        text = ""
        for sheet in wb.sheetnames:
            ws = wb[sheet]
            text += f"Sheet: {sheet}\n"
            for row in ws.iter_rows(values_only=True):
                row_text = " | ".join(str(cell) for cell in row if cell is not None)
                if row_text.strip():
                    text += row_text + "\n"
        return text.strip()

    else:
        return ""


class DocumentQueryResponse(BaseModel):
    answer: str
    sources: list
    warning: str = ""


@router.post("/ask-document", response_model=DocumentQueryResponse)
@limiter.limit("10/minute")
async def ask_about_document(
    request: Request,
    question: str,
    file: UploadFile = File(...),
    authorization: str = Header(None)
):

    if not authorization or not authorization.startswith("Bearer "):
        raise HTTPException(status_code=401, detail="Missing or invalid authorization token")

    id_token = authorization.split("Bearer ")[1]

    try:
        firebase_auth.verify_id_token(id_token)
    except Exception:
        raise HTTPException(status_code=401, detail="Invalid or expired token")

    file_bytes = await file.read()
    document_text = extract_text_from_upload(file.filename, file_bytes)

    if not document_text:
        raise HTTPException(status_code=400, detail="Could not extract text from this file. Supported formats: PDF, DOCX, XLSX.")

    truncated_doc = " ".join(document_text.split()[:3000])

    query_vector = get_embedding(question)
    results = index.query(vector=query_vector, top_k=3, include_metadata=True)
    matches = results.get("matches", [])

    law_context = ""
    sources = []
    for match in matches:
        if match["score"] >= 0.5:
            metadata = match.get("metadata", {})
            law_context += f"\n---\nSource: {metadata.get('source', 'Unknown')}\n"
            law_context += metadata.get("text", "") + "\n"
            sources.append({
                "source": metadata.get("source", "Unknown"),
                "score": round(match["score"], 4),
                "chunk_index": metadata.get("chunk_index", 0)
            })

    prompt = (
        "You are LankaTaxMind, a senior tax partner AI assistant with deep expertise "
        "in Sri Lankan tax law. The user has uploaded a specific document and wants "
        "your analysis of it in relation to Sri Lankan tax law.\n\n"
        "STRICT RULES:\n"
        "1. Base your analysis primarily on the uploaded document's actual content.\n"
        "2. Where relevant, cross-reference the uploaded document against the tax law context provided below.\n"
        "3. Cite specific clauses, figures, or sections from the uploaded document when referencing it.\n"
        "4. Cite specific sections of law when referencing the tax law context.\n"
        "5. If the document raises a tax question not covered by the provided law context, say so clearly rather than guessing.\n"
        "6. Structure your answer professionally with clear headings.\n\n"
        f"UPLOADED DOCUMENT CONTENT:\n{truncated_doc}\n\n"
        f"RELEVANT TAX LAW CONTEXT:\n{law_context if law_context else 'No directly relevant law context found in knowledge base.'}\n\n"
        f"USER QUESTION: {question}\n\n"
        "Provide your professional analysis:"
    )

    try:
        response, model_used = generate_with_retry(contents=prompt)
        return DocumentQueryResponse(
            answer=response.text.strip(),
            sources=sources
        )
    except Exception:
        return DocumentQueryResponse(
            answer="LankaTaxMind is currently experiencing high demand. Please try again shortly.",
            sources=[],
            warning="Service temporarily unavailable."
        )


class DraftRequest(BaseModel):
    document_type: str
    facts: str
    history: list = []


@router.post("/draft")
@limiter.limit("5/minute")
async def draft_document(request: Request, body: DraftRequest, authorization: str = Header(None)):

    if not authorization or not authorization.startswith("Bearer "):
        raise HTTPException(status_code=401, detail="Missing or invalid authorization token")

    id_token = authorization.split("Bearer ")[1]

    try:
        firebase_auth.verify_id_token(id_token)
    except Exception:
        raise HTTPException(status_code=401, detail="Invalid or expired token")

    query_vector = get_embedding(f"{body.document_type} {body.facts}")
    results = index.query(vector=query_vector, top_k=10, include_metadata=True)
    matches = results.get("matches", [])

    law_context = ""
    context_found = False
    for match in matches:
        if match["score"] >= 0.45:
            context_found = True
            metadata = match.get("metadata", {})
            law_context += f"\n---\nSource: {metadata.get('source', 'Unknown')}\n"
            law_context += metadata.get("text", "") + "\n"

    prompt = (
        "You are LankaTaxMind, a senior tax partner AI assistant. Draft a professional "
        f"{body.document_type} for a Sri Lankan tax matter.\n\n"
        "STRICT RULES:\n"
        "1. Use formal legal document structure and language appropriate for submission to the relevant authority.\n"
        "2. ONLY cite statutory provisions, sections, or case law that appear verbatim in the CONTEXT below. "
        "Do NOT cite, name, or reference any case, judgment, or provision that is not explicitly present in the context, "
        "even if you recall it from general knowledge. If the context does not contain supporting case law for a ground, "
        "argue that ground using only the statutory provisions available, without inventing precedent.\n"
        "3. Include placeholder fields in [SQUARE BRACKETS] for information not provided (e.g. [CLIENT NAME], [DATE], [REFERENCE NUMBER]).\n"
        "4. Structure with a clear heading, reference line, salutation, numbered grounds/arguments, and closing.\n"
        "5. Do not fabricate facts beyond what is provided — use placeholders instead.\n\n"
        f"CONTEXT (the ONLY permitted source for citations):\n"
        f"{law_context if context_found else 'No specific statutory or case law provisions were found in the knowledge base for this matter. Draft using general professional structure and clearly note that specific legal citations require manual research and verification.'}\n\n"
        f"FACTS AND INSTRUCTIONS PROVIDED BY USER:\n{body.facts}\n\n"
        f"Draft the complete {body.document_type} now, strictly following rule 2 above:"
    )

    try:
        response, _ = generate_with_retry(contents=prompt)
        draft_text = response.text.strip()
    except Exception:
        raise HTTPException(status_code=503, detail="LankaTaxMind is currently experiencing high demand. Please try again shortly.")

    doc = DocxWriter()
    doc.add_heading(body.document_type, level=1)
    for paragraph in draft_text.split("\n"):
        if paragraph.strip():
            doc.add_paragraph(paragraph.strip())

    doc.add_paragraph("")
    footer = doc.add_paragraph()
    footer_run = footer.add_run(
        "⚠ AI-GENERATED DRAFT — This document was generated by LankaTaxMind based on the firm's indexed "
        "legal database. All statutory citations, case law references, and factual placeholders must be "
        "independently verified by a qualified professional before use or submission. This draft does not "
        "constitute legal advice."
    )
    footer_run.italic = True

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename={body.document_type.replace(' ', '_')}.docx"}
    )